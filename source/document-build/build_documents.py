"""Build Cypress Command editable Office templates with python-docx.

Run with the Codex bundled Python from the workspace root. No external writes.
"""
from pathlib import Path
from datetime import datetime, timezone
from zipfile import ZipFile, ZIP_DEFLATED
import json
import uuid
import struct
from lxml import etree
from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / 'outputs' / 'Cypress_Command_Brand_System_v1.0' / 'documents'
WORK = ROOT / 'work' / 'cypress-release-docs'
LOGO = ROOT / 'outputs' / 'Cypress_Command_04C_Adoption_Kit' / 'logos' / 'png' / 'cc-04c-horizontal-primary.png'
CY, MO, AM, CH, BO = '1E4D3A', '2F6B4E', 'D97706', '0A1F16', 'F3EDE0'
BODY_FONT, DISPLAY_FONT = 'Cypress Inter', 'Cypress Fraunces'
FONT_DIR=OUT.parent/'fonts'/'desktop'

def el(tag, **attrs):
    item = OxmlElement(tag)
    for key, value in attrs.items(): item.set(qn(key), str(value))
    return item

def style_font(style, family, size, color=CH, bold=False):
    style.font.name = family
    style.font.size = Pt(size)
    style.font.color.rgb = RGBColor.from_string(color)
    style.font.bold = bold
    style.font.italic = False
    rf = style.element.get_or_add_rPr().get_or_add_rFonts()
    for key in ['ascii','hAnsi','eastAsia','cs']: rf.set(qn('w:'+key), family)
    for key in ['asciiTheme','hAnsiTheme','eastAsiaTheme','cstheme']:
        rf.attrib.pop(qn('w:'+key), None)

def field(paragraph, instruction, display='1'):
    r = paragraph.add_run()
    r._r.append(el('w:fldChar', **{'w:fldCharType':'begin'}))
    ins = el('w:instrText'); ins.set(qn('xml:space'),'preserve'); ins.text = ' '+instruction+' '
    r._r.append(ins)
    r._r.append(el('w:fldChar', **{'w:fldCharType':'separate'}))
    t=el('w:t'); t.text=display; r._r.append(t)
    r._r.append(el('w:fldChar', **{'w:fldCharType':'end'}))

def base(title, short_label):
    d = Document()
    sec=d.sections[0]
    sec.page_width=Inches(8.5); sec.page_height=Inches(11)
    sec.top_margin=Inches(1.66); sec.bottom_margin=Inches(.8)
    sec.left_margin=Inches(.8); sec.right_margin=Inches(.8)
    sec.header_distance=Inches(.36); sec.footer_distance=Inches(.38)
    styles=d.styles
    normal=styles['Normal']; style_font(normal, BODY_FONT, 11.5)
    normal.paragraph_format.line_spacing=1.18
    normal.paragraph_format.space_after=Pt(8)
    normal.paragraph_format.widow_control=True
    for name, size, family in [('Title',28,DISPLAY_FONT),('Subtitle',13,BODY_FONT),('Heading 1',19,DISPLAY_FONT),('Heading 2',12,BODY_FONT)]:
        st=styles[name]; style_font(st,family,size,'000000', name=='Heading 2')
        st.paragraph_format.space_before=Pt(16 if name.startswith('Heading') else 0)
        st.paragraph_format.space_after=Pt(9)
        st.paragraph_format.keep_with_next=True
        st.paragraph_format.line_spacing=1.1
        ppr=st.element.get_or_add_pPr()
        for node in list(ppr):
            if node.tag==qn('w:pBdr'): ppr.remove(node)
    for name,size,color in [('CC Metadata',9.5,MO),('CC Footer',8.5,CH),('CC Small',9.5,CH),('CC Table',10.5,CH)]:
        st=styles.add_style(name,WD_STYLE_TYPE.PARAGRAPH); st.base_style=normal
        style_font(st,BODY_FONT,size,color)
        st.paragraph_format.space_after=Pt(5)
        st.paragraph_format.line_spacing=1.12
    st=styles.add_style('CC Placeholder',WD_STYLE_TYPE.CHARACTER)
    style_font(st,BODY_FONT,11.5,MO)
    for name in ['Header','Footer']:
        style_font(styles[name],BODY_FONT,9,'000000')
    hp=sec.header.paragraphs[0]
    hp.paragraph_format.space_after=Pt(0)
    shape=hp.add_run().add_picture(str(LOGO),width=Inches(2.6))
    shape._inline.docPr.set('descr','Cypress Command approved rounded 04C logo')
    hp.paragraph_format.left_indent=Inches(-.17)
    foot=sec.footer.paragraphs[0]; foot.style=styles['CC Footer']
    foot.paragraph_format.tab_stops.add_tab_stop(Inches(6.9),WD_TAB_ALIGNMENT.RIGHT)
    foot.add_run('CYPRESS COMMAND  |  532 Alonda Drive, Lafayette, LA 70503  |  '+short_label+'\t')
    field(foot,'PAGE'); foot.add_run(' / '); field(foot,'NUMPAGES')
    s=d.settings.element
    s.append(el('w:updateFields', **{'w:val':'true'}))
    s.append(el('w:embedTrueTypeFonts', **{'w:val':'true'}))
    s.append(el('w:saveSubsetFonts', **{'w:val':'false'}))
    cp=d.core_properties
    cp.title=title; cp.subject='Cypress Command editable brand template'
    cp.author='Cypress Command'; cp.last_modified_by='Cypress Command'
    cp.keywords='Cypress Command, 04C, template, v1.0'
    cp.comments=''; cp.created=datetime(2026,9,6,tzinfo=timezone.utc); cp.modified=cp.created
    return d

def p(d,text='',style=None):
    return d.add_paragraph(text,style)

def placeholder(d,text,space_after=8):
    a=p(d); a.add_run(text,'CC Placeholder'); a.paragraph_format.space_after=Pt(space_after)
    return a

def h(d,text): return d.add_heading(text,level=1)

def table(d,headers,rows,widths):
    t=d.add_table(rows=1,cols=len(headers)); t.alignment=WD_TABLE_ALIGNMENT.LEFT
    t.autofit=False
    for col,w in zip(t.columns,widths): col.width=Inches(w)
    pr=t._tbl.tblPr
    borders=el('w:tblBorders')
    for edge in ['top','left','bottom','right','insideH','insideV']:
        borders.append(el('w:'+edge, **{'w:val':'single','w:sz':'4','w:color':'D9D9D9'}))
    pr.append(borders)
    margins=el('w:tblCellMar')
    for edge in ['top','bottom']: margins.append(el('w:'+edge,**{'w:w':'110','w:type':'dxa'}))
    for edge in ['left','right']: margins.append(el('w:'+edge,**{'w:w':'125','w:type':'dxa'}))
    pr.append(margins)
    for idx,text in enumerate(headers):
        c=t.rows[0].cells[idx]; c.width=Inches(widths[idx]); c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
        c._tc.get_or_add_tcPr().append(el('w:shd', **{'w:fill':CY}))
        pa=c.paragraphs[0]; pa.style=d.styles['CC Table']; pa.paragraph_format.space_after=Pt(0)
        rr=pa.add_run(text); rr.bold=True; rr.font.color.rgb=RGBColor(255,255,255)
    t.rows[0]._tr.get_or_add_trPr().append(el('w:tblHeader'))
    for ri,row in enumerate(rows):
        cells=t.add_row().cells
        for ci,text in enumerate(row):
            c=cells[ci]; c.width=Inches(widths[ci]); c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if ri%2==1: c._tc.get_or_add_tcPr().append(el('w:shd', **{'w:fill':'F6F6F3'}))
            pa=c.paragraphs[0]; pa.style=d.styles['CC Table']; pa.paragraph_format.space_after=Pt(0)
            pa.add_run(text)
        trpr=cells[0]._tc.getparent().get_or_add_trPr(); trpr.append(el('w:cantSplit'))
    p(d).paragraph_format.space_after=Pt(0)
    return t

def newpage(d,label):
    d.add_page_break()
    p(d,label.upper(),'CC Metadata')

def embed_fonts(contents):
    ns={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main','r':'http://schemas.openxmlformats.org/officeDocument/2006/relationships'}
    table_root=etree.fromstring(contents['word/fontTable.xml'])
    relns='http://schemas.openxmlformats.org/package/2006/relationships'
    relroot=etree.Element('{'+relns+'}Relationships', nsmap={None:relns})
    types=etree.fromstring(contents['[Content_Types].xml'])
    ctns='http://schemas.openxmlformats.org/package/2006/content-types'
    etree.SubElement(types,'{'+ctns+'}Default',Extension='odttf',ContentType='application/vnd.openxmlformats-officedocument.obfuscatedFont')
    for family,sources in [(BODY_FONT,[('Regular','CypressInter-Regular.ttf'),('Bold','CypressInter-Bold.ttf')]),(DISPLAY_FONT,[('Regular','CypressFraunces-Regular.ttf')])]:
        node=etree.SubElement(table_root,'{'+ns['w']+'}font',attrib={qn('w:name'):family})
        etree.SubElement(node,qn('w:family'),attrib={qn('w:val'):'swiss' if family==BODY_FONT else 'roman'})
        etree.SubElement(node,qn('w:pitch'),attrib={qn('w:val'):'variable'})
        for face,filename in sources:
            fontdata=bytearray((FONT_DIR/filename).read_bytes())
            # Enforce that these OFL sources do not request restricted embedding.
            nt=struct.unpack_from('>H',fontdata,4)[0]
            for ti in range(nt):
                off=12+16*ti
                if fontdata[off:off+4]==b'OS/2':
                    pos=struct.unpack_from('>I',fontdata,off+8)[0]
                    assert struct.unpack_from('>H',fontdata,pos+8)[0] & 2 == 0, 'Restricted font embedding'
            key=uuid.uuid5(uuid.NAMESPACE_URL,'cypress-command/v1.0/'+filename)
            keybytes=bytes.fromhex(key.hex)[::-1]
            for ii in range(32): fontdata[ii]^=keybytes[ii%16]
            target='fonts/'+filename.replace('.ttf','.odttf')
            contents['word/'+target]=bytes(fontdata)
            rid='rId'+filename.replace('-','').replace('.ttf','')
            etree.SubElement(relroot,'{'+relns+'}Relationship',Id=rid,Type=ns['r']+'/font',Target=target)
            etree.SubElement(node,qn('w:embed'+face),attrib={qn('r:id'):rid,qn('w:fontKey'):'{'+str(key).upper()+'}',qn('w:subsetted'):'false'})
    for member,root in [('word/fontTable.xml',table_root),('word/_rels/fontTable.xml.rels',relroot),('[Content_Types].xml',types)]:
        contents[member]=etree.tostring(root,xml_declaration=True,encoding='UTF-8',standalone=True)
    return contents

def save(d,name):
    path=OUT/(name+'.docx'); d.save(path)
    with ZipFile(path) as z: contents={x:z.read(x) for x in z.namelist()}
    contents=embed_fonts(contents)
    with ZipFile(path,'w',ZIP_DEFLATED) as z:
        for n,data in contents.items(): z.writestr(n,data)
    contents['[Content_Types].xml']=contents['[Content_Types].xml'].replace(b'application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml',b'application/vnd.openxmlformats-officedocument.wordprocessingml.template.main+xml')
    with ZipFile(OUT/(name+'.dotx'),'w',ZIP_DEFLATED) as z:
        for n,data in contents.items(): z.writestr(n,data)
    print(path)

def letter():
    d=base('Business correspondence','cypresscommand.com')
    p(d,'[Date]','CC Metadata')
    p(d,'[Recipient name]\n[Recipient role]\n[Organization]\n[Mailing address]')
    p(d,'Business correspondence','Title')
    placeholder(d,'[Subject of this letter]')
    p(d,'Dear [Recipient name],')
    placeholder(d,'[State the purpose of this letter and the action or decision requested.]',15)
    placeholder(d,'[Add the relevant facts, supporting context, and any dates the recipient needs.]',15)
    placeholder(d,'[Close with the next step, responsible person, and response date.]',22)
    p(d,'Sincerely,')
    p(d,'Adam Abdalla\nOwner / Founder\n[Email]  |  [Phone]')
    p(d,'[Attachment reference if applicable]','CC Small')
    save(d,'Cypress_Command_Letterhead')

def proposal():
    d=base('Project proposal','PROJECT PROPOSAL')
    p(d,'PROJECT [Reference]  |  [Date]  |  VERSION [Number]','CC Metadata')
    p(d,'Project proposal','Title')
    p(d,'[Project name]','Subtitle')
    p(d,'Prepared for [Client organization]\nPrepared by Adam Abdalla, Owner / Founder\n[Email]  |  [Phone]','CC Small')
    h(d,'Recommended engagement')
    placeholder(d,'[Describe the recommended work, the business need it addresses, and the decision requested.]')
    h(d,'Intended outcome')
    placeholder(d,'[Describe the practical result the client should receive. State how completion will be evaluated without promising unsupported results.]')
    h(d,'Context and constraints')
    placeholder(d,'[Summarize the current position, relevant deadlines, dependencies, and known constraints.]')
    h(d,'Success measures')
    table(d,['Measure','Completion criterion','Evidence'],[
        ['[Measure]','[Criterion]','[Source or artifact]'],
        ['[Measure]','[Criterion]','[Source or artifact]'],
    ],[1.7,2.7,2.5])
    newpage(d,'Scope of work')
    h(d,'Deliverables')
    p(d,'[Describe the proposed scope. Add or remove rows to match the engagement.]')
    table(d,['Deliverable','Included work','Review point'],[
        ['[Deliverable]','[Activities and output]','[Owner and review]'],
        ['[Deliverable]','[Activities and output]','[Owner and review]'],
        ['[Deliverable]','[Activities and output]','[Owner and review]'],
    ],[1.6,3.5,1.8])
    h(d,'Client inputs and responsibilities')
    placeholder(d,'[Identify the access, documents, decisions, and staff participation required from the client. Assign an owner and due date for each material dependency.]')
    h(d,'Exclusions')
    placeholder(d,'[State the work, deliverables, third party costs, and ongoing services that are outside this scope.]')
    h(d,'Scope changes')
    placeholder(d,'[Describe how requested changes will be evaluated, priced, and approved before work is added.]')
    newpage(d,'Fees and next steps')
    h(d,'Proposed fees')
    table(d,['Item','Basis','Amount'],[
        ['[Service or phase]','[Fixed fee or rate]','[Currency and amount]'],
        ['[Expenses or third party costs]','[Included or additional]','[Currency and amount]'],
        ['Total proposed fee','[Tax treatment]','[Currency and amount]'],
    ],[2.75,2.25,1.9])
    p(d,'[Payment timing, billing cadence, deposit if any, and proposal validity date.]','CC Small')
    h(d,'Proposed schedule')
    table(d,['Milestone','Target','Dependency'],[
        ['[Milestone]','[Date or duration]','[Required input]'],
        ['[Milestone]','[Date or duration]','[Required input]'],
    ],[2.4,1.9,2.6])
    h(d,'Commercial assumptions')
    placeholder(d,'[Document assumptions that affect price or timing, including access, availability, location, and travel.]')
    h(d,'Next step')
    placeholder(d,'[Name the decision maker, requested decision, and date for confirming the proposed scope and commercial terms.]')
    p(d,'This proposal describes a prospective engagement. Any binding services, payment, liability, and other legal terms require a separate approved agreement.','CC Small')
    save(d,'Cypress_Command_Proposal')

def report():
    d=base('Operating report','OPERATING REPORT')
    p(d,'[Entity or portfolio]  |  PERIOD [Start date to end date]','CC Metadata')
    p(d,'Operating report','Title')
    p(d,'[Business unit or property name]','Subtitle')
    p(d,'Prepared by Adam Abdalla, Owner / Founder  |  Issued [Date]\nReporting cutoff [Date and time]  |  Version [Number]','CC Small')
    h(d,'Management summary')
    placeholder(d,'[State the overall position, the material change during this period, and any decision or intervention needed.]')
    h(d,'Performance measures')
    table(d,['Measure','Current','Comparator','Source'],[
        ['[Metric and unit]','[Value]','[Target or prior]','[Record and date]'],
        ['[Metric and unit]','[Value]','[Target or prior]','[Record and date]'],
        ['[Metric and unit]','[Value]','[Target or prior]','[Record and date]'],
        ['[Metric and unit]','[Value]','[Target or prior]','[Record and date]'],
    ],[2.05,1.1,1.7,2.05])
    p(d,'[Define each measure and comparator. Distinguish zero from unavailable data, and identify estimates.]','CC Small')
    h(d,'Material changes')
    placeholder(d,'[Explain significant variances and changes, their causes, and the evidence supporting the explanation.]')
    newpage(d,'Work and risks')
    h(d,'Workstream status')
    table(d,['Workstream','Progress and next action','Owner','Due'],[
        ['[Workstream]','[Status and action]','[Name]','[Date]'],
        ['[Workstream]','[Status and action]','[Name]','[Date]'],
        ['[Workstream]','[Status and action]','[Name]','[Date]'],
    ],[1.6,2.9,1.2,1.2])
    h(d,'Open risks and issues')
    table(d,['Risk or issue','Impact and response','Owner','Review'],[
        ['[Description]','[Impact and mitigation]','[Name]','[Date]'],
        ['[Description]','[Impact and mitigation]','[Name]','[Date]'],
    ],[2,2.5,1.2,1.2])
    h(d,'Dependencies and escalation')
    placeholder(d,'[Identify unresolved dependencies, the person who can resolve them, and when escalation is required.]')
    h(d,'Completed work')
    placeholder(d,'[Record the material items closed this period and link to acceptance or completion evidence.]')
    newpage(d,'Decisions and next period')
    h(d,'Decisions requested')
    table(d,['Decision','Recommendation','Decision maker','Needed by'],[
        ['[Decision]','[Recommendation and basis]','[Name]','[Date]'],
        ['[Decision]','[Recommendation and basis]','[Name]','[Date]'],
    ],[1.6,2.65,1.45,1.2])
    h(d,'Priorities for the next period')
    table(d,['Priority','Expected output','Owner','Due'],[
        ['[Priority]','[Specific deliverable]','[Name]','[Date]'],
        ['[Priority]','[Specific deliverable]','[Name]','[Date]'],
        ['[Priority]','[Specific deliverable]','[Name]','[Date]'],
    ],[1.6,2.9,1.2,1.2])
    h(d,'Source record')
    table(d,['Source','As of','Coverage or limitation'],[
        ['[File or system record]','[Date]','[Coverage and limitations]'],
        ['[File or system record]','[Date]','[Coverage and limitations]'],
    ],[2.4,1.25,3.25])
    p(d,'[Reviewer name and role]  |  Reviewed [Date]  |  Next report [Date]','CC Small')
    save(d,'Cypress_Command_Operating_Report')

if __name__=='__main__':
    OUT.mkdir(parents=True,exist_ok=True); WORK.mkdir(parents=True,exist_ok=True)
    letter(); proposal(); report()
